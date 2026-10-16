"""
Modul monitor_mitratel_default untuk backend Django.
Mendukung delegasi otomatis ke Worker_Verticality/monitor_mitratel_default.py jika tersedia,
atau menjalankan logika monitoring mandiri jika Worker_Verticality belum di-mount ke container.
"""
import argparse
from datetime import datetime, timezone, timedelta
import importlib.util
import os
import re
import sys
import time
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==============================================================================
# CEK APAKAH FILE DI Worker_Verticality DAPAT DIAKSES LANGSUNG
# ==============================================================================
_candidates = [
    os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Worker_Verticality', 'monitor_mitratel_default.py'),
    os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'Worker_Verticality', 'monitor_mitratel_default.py'),
    '/app/Worker_Verticality/monitor_mitratel_default.py',
    '/Worker_Verticality/monitor_mitratel_default.py',
]

_external_mod = None
for _p in _candidates:
    if os.path.exists(_p) and os.path.abspath(_p) != os.path.abspath(__file__):
        try:
            _spec = importlib.util.spec_from_file_location("worker_mitratel_mod_live", _p)
            if _spec and _spec.loader:
                _mod = importlib.util.module_from_spec(_spec)
                _spec.loader.exec_module(_mod)
                _external_mod = _mod
                break
        except Exception:
            pass

if _external_mod is not None:
    for _k, _v in _external_mod.__dict__.items():
        if not _k.startswith("__"):
            globals()[_k] = _v
else:
    # ==============================================================================
    # FALLBACK MANDIRI JIKA Worker_Verticality BELUM DI-MOUNT
    # ==============================================================================
    DEFAULT_API_KEY = os.getenv(
        "MAIVIEW_API_KEY",
        "nms_live_0c05c2b6_d09b6897332bcc9fde952de761b6e15ba8cb9366817b22231ca126a2b933106a",
    )
    API_URL = os.getenv("MAIVIEW_API_URL", "https://api.maiviewmitratel.id/v1/sites")
    DEFAULT_BRANCH = os.getenv("MONITOR_BRANCH", "Mitratel Default")
    DEFAULT_EXCEL_FILE = os.getenv(
        "EXCEL_OUTPUT",
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "monitoring_mitratel_hourly.xlsx"),
    )
    MEDIA_REPORTS_DIR = os.getenv("MEDIA_REPORTS_DIR", "")

    EXTRA_SITE_IDS = [
        "bd3deb2d-6e81-5540-9eed-04fd397d5ca6",  # dmt new ikn (DMT_NEW_IKN)
        "83a838ed-698f-414a-a37c-5669f5a1e4b6",  # dmt isat tamangapa (DMT ISAT_TAMANGAPA)
        "347e50b4-eabb-46f0-a96d-a9da03934119",  # margodadii (MARGODADI)
        "575e8551-5e53-5c3d-abb2-39984e5cd8db",  # pelab ajibata (PELAB AJIBATA)
        "c6dc7fe4-8d26-5587-9101-1ead5d9faf65",  # dagang kelambir (Dagang Kelambir)
        "ea5ae5e7-a752-501a-99db-b0738bf9daa2",  # permanenisasi (PERMANENISASICIOKONGSUKARESIKSIDAMULIH)
        "661cc75b-37ec-539d-9197-853f91dd7f16",  # jalanmustikasari (JALANMUSTIKASARI)
        "ee11245f-deb1-43fe-abc1-d5d52d399053",  # nayakaws
    ]

    WIB = timezone(timedelta(hours=7))

    def clean_auth_token(token: str) -> str:
        if not token:
            return ""
        return re.sub(r"(?i)^bearer\s+", "", token).strip()

    def get_seconds_until_next_hour() -> float:
        now = datetime.now(WIB)
        next_hour = (now + timedelta(hours=1)).replace(minute=0, second=0, microsecond=0)
        diff = (next_hour - now).total_seconds()
        return max(5.0, diff + 1.5)

    def fetch_sites_by_branch(api_key: str, branch_name: str = DEFAULT_BRANCH) -> tuple[list[dict], datetime]:
        clean_key = clean_auth_token(api_key)
        headers = {
            "Accept": "application/json",
            "Authorization": f"Bearer {clean_key}",
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) MaiView-Monitor/1.0",
        }

        waktu_tarik = datetime.now(WIB)
        all_sites = []
        page = 1
        per_page = 100
        max_retries = 3

        while True:
            params = {"page": page, "per_page": per_page, "is_active": "true"}
            success = False

            for attempt in range(1, max_retries + 1):
                try:
                    res = requests.get(API_URL, headers=headers, params=params, timeout=25)
                    if res.status_code == 200:
                        data = res.json()
                        items = data.get("data") or []
                        all_sites.extend(items)
                        success = True
                        if len(items) < per_page:
                            break
                        page += 1
                        break
                    else:
                        print(f"[!] Gagal halaman {page}: HTTP {res.status_code} (Percobaan {attempt}/{max_retries})")
                        if attempt < max_retries:
                            time.sleep(3 * attempt)
                except requests.RequestException as e:
                    print(f"[!] Error koneksi API halaman {page}: {e} (Percobaan {attempt}/{max_retries})")
                    if attempt < max_retries:
                        time.sleep(3 * attempt)

            if not success or (success and len(items) < per_page):
                break

        target_clean = branch_name.strip().lower()
        extra_set = set(EXTRA_SITE_IDS)
        filtered_sites = [
            s for s in all_sites
            if (
                str(s.get("branch_name", "")).strip().lower() == target_clean
                or s.get("id") in extra_set
            )
        ]
        return filtered_sites, waktu_tarik

    def update_hourly_excel_matrix(
        sites: list[dict],
        waktu_tarik: datetime,
        file_path: str = DEFAULT_EXCEL_FILE,
        align_hour: bool = True,
        media_dir: str = MEDIA_REPORTS_DIR,
    ) -> str:
        sheet_name = waktu_tarik.strftime("%Y-%m-%d")
        hour_header = waktu_tarik.strftime("%H:00") if align_hour else waktu_tarik.strftime("%H:%M")

        dir_path = os.path.dirname(os.path.abspath(file_path))
        if dir_path:
            os.makedirs(dir_path, exist_ok=True)

        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_bold = Font(name="Calibri", size=10, bold=True)
        font_regular = Font(name="Calibri", size=10)
        
        fill_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        fill_header_hour = PatternFill(start_color="2C5E8A", end_color="2C5E8A", fill_type="solid")
        fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

        fill_online = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
        font_online = Font(name="Calibri", size=10, bold=True, color="2E7D32")

        fill_offline = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
        font_offline = Font(name="Calibri", size=10, bold=True, color="C62828")

        thin_border_side = Side(border_style="thin", color="D1D5DB")
        thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

        if os.path.exists(file_path):
            try:
                wb = openpyxl.load_workbook(file_path)
            except Exception:
                wb = openpyxl.Workbook()
        else:
            wb = openpyxl.Workbook()

        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            if len(wb.sheetnames) == 1 and wb.sheetnames[0] == "Sheet":
                ws = wb.active
                ws.title = sheet_name
            else:
                ws = wb.create_sheet(title=sheet_name)

        ws.views.sheetView[0].showGridLines = True
        is_new_sheet = (ws.max_row < 2 or ws.cell(row=1, column=1).value is None)

        api_site_map = {}
        for s in sites:
            scode = str(s.get("site_code", "")).strip()
            sname = str(s.get("name", "")).strip()
            sstatus = str(s.get("status", "-")).upper()
            api_site_map[scode] = sstatus
            api_site_map[sname] = sstatus

        if is_new_sheet:
            ws.cell(row=1, column=1, value="Nama Site").font = font_header
            ws.cell(row=1, column=1).fill = fill_header
            ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=1, column=1).border = thin_border

            ws.cell(row=1, column=2, value="Site Code").font = font_header
            ws.cell(row=1, column=2).fill = fill_header
            ws.cell(row=1, column=2).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=1, column=2).border = thin_border

            sorted_sites = sorted(sites, key=lambda x: str(x.get("name", "")).lower())

            for idx, s in enumerate(sorted_sites, start=2):
                cell_name = ws.cell(row=idx, column=1, value=s.get("name", "-"))
                cell_code = ws.cell(row=idx, column=2, value=s.get("site_code", "-"))
                
                cell_name.font = font_bold
                cell_name.border = thin_border
                cell_name.alignment = Alignment(horizontal="left", vertical="center")

                cell_code.font = font_regular
                cell_code.border = thin_border
                cell_code.alignment = Alignment(horizontal="center", vertical="center")

                if idx % 2 == 0:
                    cell_name.fill = fill_zebra
                    cell_code.fill = fill_zebra

            target_col = 3
        else:
            existing_site_names = set()
            existing_site_codes = set()
            for r in range(2, ws.max_row + 1):
                s_name_val = str(ws.cell(row=r, column=1).value or "").strip().lower()
                s_code_val = str(ws.cell(row=r, column=2).value or "").strip().lower()
                if s_name_val:
                    existing_site_names.add(s_name_val)
                if s_code_val:
                    existing_site_codes.add(s_code_val)

            for s in sites:
                s_name = str(s.get("name", "")).strip()
                s_code = str(s.get("site_code", "")).strip()
                if s_name.lower() not in existing_site_names and s_code.lower() not in existing_site_codes:
                    new_row = ws.max_row + 1
                    cell_name = ws.cell(row=new_row, column=1, value=s_name)
                    cell_code = ws.cell(row=new_row, column=2, value=s_code)
                    cell_name.font = font_bold
                    cell_name.border = thin_border
                    cell_name.alignment = Alignment(horizontal="left", vertical="center")
                    cell_code.font = font_regular
                    cell_code.border = thin_border
                    cell_code.alignment = Alignment(horizontal="center", vertical="center")
                    if new_row % 2 == 0:
                        cell_name.fill = fill_zebra
                        cell_code.fill = fill_zebra
                    existing_site_names.add(s_name.lower())
                    existing_site_codes.add(s_code.lower())

            target_col = None
            for col_idx in range(3, ws.max_column + 1):
                col_val = str(ws.cell(row=1, column=col_idx).value or "").strip()
                if col_val == hour_header:
                    target_col = col_idx
                    break
            
            if target_col is None:
                target_col = ws.max_column + 1

        header_cell = ws.cell(row=1, column=target_col, value=hour_header)
        header_cell.font = font_header
        header_cell.fill = fill_header_hour
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cell.border = thin_border
        ws.row_dimensions[1].height = 26

        for row_idx in range(2, ws.max_row + 1):
            site_name_val = str(ws.cell(row=row_idx, column=1).value or "").strip()
            site_code_val = str(ws.cell(row=row_idx, column=2).value or "").strip()

            status = api_site_map.get(site_code_val) or api_site_map.get(site_name_val) or "-"
            
            cell = ws.cell(row=row_idx, column=target_col, value=status)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            ws.row_dimensions[row_idx].height = 20

            if status == "ONLINE":
                cell.fill = fill_online
                cell.font = font_online
            elif status == "OFFLINE":
                cell.fill = fill_offline
                cell.font = font_offline
            else:
                cell.font = font_regular

        ws.freeze_panes = "C2"
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 16
        for col_i in range(3, ws.max_column + 1):
            col_letter = get_column_letter(col_i)
            ws.column_dimensions[col_letter].width = 12

        saved = False
        for attempt in range(1, 4):
            try:
                wb.save(file_path)
                saved = True
                break
            except PermissionError:
                print(f"[!] PERINGATAN: File '{file_path}' terkunci program lain. Mencoba ulang dalam 5 detik...")
                time.sleep(5)
            except Exception as e:
                print(f"[!] Gagal menyimpan file Excel: {e}")
                break

        if saved and media_dir:
            try:
                os.makedirs(media_dir, exist_ok=True)
                media_file = os.path.join(media_dir, os.path.basename(file_path))
                wb.save(media_file)
                print(f"[+] Salinan berhasil disimpan ke media directory: {media_file}")
            except Exception as e:
                print(f"[!] Gagal menyalin ke media directory: {e}")

        return os.path.abspath(file_path)

    def run_single_monitoring(
        api_key: str = DEFAULT_API_KEY,
        branch_name: str = DEFAULT_BRANCH,
        file_path: str = DEFAULT_EXCEL_FILE,
        align_hour: bool = True,
        media_dir: str = MEDIA_REPORTS_DIR,
    ) -> str | None:
        sites, waktu_tarik = fetch_sites_by_branch(api_key=api_key, branch_name=branch_name)
        if not sites:
            print(f"[!] Tidak ada site yang ditemukan untuk branch '{branch_name}'.")
            return None

        saved_file = update_hourly_excel_matrix(
            sites=sites,
            waktu_tarik=waktu_tarik,
            file_path=file_path,
            align_hour=align_hour,
            media_dir=media_dir,
        )

        online_cnt = sum(1 for s in sites if str(s.get("status", "")).lower() == "online")
        offline_cnt = sum(1 for s in sites if str(s.get("status", "")).lower() == "offline")
        jam_str = waktu_tarik.strftime("%H:00 WIB") if align_hour else waktu_tarik.strftime("%H:%M WIB")

        print(f"[{waktu_tarik.strftime('%Y-%m-%d %H:%M:%S')} WIB] Berhasil update kolom jam '{jam_str}'")
        print(f"  -> Total: {len(sites)} site | Online: {online_cnt} | Offline: {offline_cnt}")
        print(f"  -> File Excel: {saved_file}\n")
        return saved_file
